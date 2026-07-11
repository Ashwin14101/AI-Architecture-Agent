#!/usr/bin/env python3
"""
Generate AI Architecture Builder SDD as .docx for Google Docs editing.
Same content as the PDF - compact, professional, black & white.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AI Architecture Builder.docx")


def set_cell_shading(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Arial'
        set_cell_shading(cell, 'EEEEEE')

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9)
            run.font.name = 'Arial'

    return table


def add_folder_structure(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = 'Arial'


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.size = Pt(10)
        run_b.font.name = 'Arial'
        run_n = p.add_run(text)
        run_n.font.size = Pt(10)
        run_n.font.name = 'Arial'
    else:
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.name = 'Arial'


def heading1(doc, text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = 'Arial'


def heading2(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = 'Arial'
        run.font.size = Pt(13)


def build():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # ===== TITLE =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('AI Architecture Builder')
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0, 0, 0)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run('Software Design Document  •  Version 1.0  •  July 2026')
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(85, 85, 85)

    add_hr(doc)

    # ===== PROJECT OVERVIEW =====
    heading1(doc, 'Project Overview')
    body(doc,
        'The AI Architecture Builder is an enterprise platform that uses a multi-agent AI system to '
        'automatically generate complete cloud architectures from requirement documents. Users upload '
        'project requirements, and the system produces system diagrams, database schemas, API designs, '
        'Terraform scripts, security validations, and cost estimates — all without manual architecture work.'
    )

    heading2(doc, 'Objectives')
    bullet(doc, 'Automate end-to-end cloud architecture generation from requirement documents.')
    bullet(doc, 'Use 16 specialized AI agents that collaborate to produce production-ready output.')
    bullet(doc, 'Generate deployable Terraform code, database schemas, and API specifications.')
    bullet(doc, 'Provide AI-powered architecture review with quality scoring.')
    bullet(doc, 'Support interactive refinement through AI chat.')

    heading2(doc, 'Goals')
    bullet(doc, ' Reduce architecture design time from weeks to hours. Standardize architecture practices across teams.', 'Business:')
    bullet(doc, ' Scalable multi-agent orchestration. Cloud-agnostic (AWS, Azure, GCP). Extensible agent framework.', 'Technical:')
    bullet(doc, ' Context-aware decisions using RAG. Explainable recommendations. Reliable validation with low false positives.', 'AI:')

    heading2(doc, 'Who Uses It')
    bullet(doc, ' — accelerate design and validate decisions.', 'Architects')
    bullet(doc, ' — generate architectures without waiting for architects.', 'Dev Teams')
    bullet(doc, ' — receive production-ready Terraform and pipelines.', 'DevOps')
    bullet(doc, ' — review architectures with AI quality scores.', 'Managers')

    add_hr(doc)

    # ===== PHASE 1 =====
    heading1(doc, 'Phase 1 — Project Initialization & Authentication')
    body(doc,
        'Set up the monorepo project structure, configure the development environment, implement user '
        'authentication with JWT, create the database schema with migrations, and build the basic frontend '
        'shell with login, registration, and an empty dashboard. This is the foundation — nothing else '
        'works without auth and database.'
    )

    heading2(doc, 'What to do')
    bullet(doc, 'Initialize monorepo with frontend (React + Vite + TypeScript) and backend (Node.js + Express + TypeScript) packages.')
    bullet(doc, 'Set up PostgreSQL database with migration tool (Prisma or Knex).')
    bullet(doc, 'Implement user registration, login, logout, password reset.')
    bullet(doc, 'JWT access tokens (15 min expiry) + refresh tokens (7 day) with rotation.')
    bullet(doc, 'Build frontend shell: login page, register page, sidebar + header layout, empty dashboard.')
    bullet(doc, 'Set up audit logging for all auth events.')
    bullet(doc, 'Configure ESLint, Prettier, Jest, CI/CD pipeline.')

    heading2(doc, 'Folder Structure')
    add_folder_structure(doc, [
        "ai-architecture-agent/",
        "├── frontend/",
        "│   ├── src/",
        "│   │   ├── components/         # Reusable UI components",
        "│   │   │   ├── common/         # Button, Input, Modal, Loader",
        "│   │   │   └── layout/         # Sidebar, Header, AppLayout",
        "│   │   ├── pages/",
        "│   │   │   ├── auth/           # LoginPage, RegisterPage, ForgotPassword",
        "│   │   │   └── dashboard/      # DashboardPage",
        "│   │   ├── services/           # API client (axios instance, auth service)",
        "│   │   ├── store/              # State management (Zustand or Redux)",
        "│   │   ├── hooks/              # useAuth, useApi custom hooks",
        "│   │   ├── utils/              # Helpers, validators, formatters",
        "│   │   ├── types/              # TypeScript interfaces",
        "│   │   ├── styles/             # Global CSS, design tokens",
        "│   │   ├── routes/             # Route config with protected routes",
        "│   │   ├── App.tsx             # Root component",
        "│   │   └── main.tsx            # Entry point",
        "│   ├── package.json",
        "│   └── vite.config.ts",
        "├── backend/",
        "│   ├── src/",
        "│   │   ├── config/             # DB, auth, app config files",
        "│   │   ├── controllers/        # auth.controller.ts",
        "│   │   ├── services/           # auth.service.ts, audit.service.ts",
        "│   │   ├── models/             # user.model.ts, organization.model.ts",
        "│   │   ├── middleware/          # auth.middleware.ts, error-handler.ts",
        "│   │   ├── routes/             # auth.routes.ts",
        "│   │   ├── migrations/         # 001_create_users.ts, 002_create_orgs.ts",
        "│   │   ├── utils/              # jwt.ts, password.ts, logger.ts",
        "│   │   ├── types/              # TypeScript interfaces",
        "│   │   └── app.ts              # Express app entry point",
        "│   └── package.json",
        "├── shared/                      # Shared types between frontend/backend",
        "├── docker-compose.yml           # PostgreSQL + Redis for local dev",
        "├── .github/workflows/           # CI/CD pipelines",
        "└── package.json                 # Root workspace config",
    ])

    add_hr(doc)

    # ===== PHASE 2 =====
    heading1(doc, 'Phase 2 — Backend API Development')
    body(doc,
        'Build all backend REST APIs for project management, document upload and processing, architecture '
        'generation triggers, version management, review handling, deployment management, and chat. Each '
        'API module has a controller (handles request/response), a service (business logic), and a model '
        '(database). WebSocket support for real-time progress updates.'
    )

    heading2(doc, 'What to do')
    bullet(doc, 'Build Project CRUD APIs — create, list, get, update, archive projects.')
    bullet(doc, 'Build Document Upload API — multipart upload, validate format (PDF/DOCX/TXT/MD), store in S3, trigger processing.')
    bullet(doc, 'Build Architecture APIs — trigger generation, get current architecture, get database schema, get API specs.')
    bullet(doc, 'Build Version APIs — list versions, get version, compare two versions, rollback.')
    bullet(doc, 'Build Review APIs — request review, list reviews, get review details, respond to findings.')
    bullet(doc, 'Build Deployment APIs — initiate, approve, list, rollback deployments.')
    bullet(doc, 'Build Chat APIs — create conversation, send message, get message history.')
    bullet(doc, 'Build Cloud Mapping and Terraform APIs — get mappings, get/download/validate Terraform.')
    bullet(doc, 'Set up WebSocket server for real-time generation progress and deployment status.')
    bullet(doc, 'Implement RBAC middleware (Viewer, Editor, Approver, Admin roles).')
    bullet(doc, 'Add rate limiting, input validation, and error handling middleware.')

    heading2(doc, 'API Reference')
    api_data = [
        ['POST /api/auth/register', 'Create user account'],
        ['POST /api/auth/login', 'Login, returns JWT tokens'],
        ['POST /api/auth/refresh', 'Refresh access token'],
        ['POST /api/auth/logout', 'Invalidate session'],
        ['POST /api/projects', 'Create new project'],
        ['GET /api/projects', 'List user projects'],
        ['GET /api/projects/:id', 'Get project details'],
        ['PUT /api/projects/:id', 'Update project settings'],
        ['DELETE /api/projects/:id', 'Archive project'],
        ['POST /api/projects/:id/documents', 'Upload requirement document'],
        ['GET /api/projects/:id/documents', 'List project documents'],
        ['POST /api/projects/:id/generate', 'Start architecture generation'],
        ['GET /api/projects/:id/architecture', 'Get current architecture'],
        ['GET /api/projects/:id/architecture/database', 'Get database schema'],
        ['GET /api/projects/:id/architecture/apis', 'Get API specification'],
        ['GET /api/projects/:id/cloud-mapping', 'Get cloud service mapping'],
        ['GET /api/projects/:id/terraform', 'Get Terraform code'],
        ['GET /api/projects/:id/terraform/download', 'Download Terraform as zip'],
        ['POST /api/projects/:id/terraform/validate', 'Validate Terraform syntax'],
        ['POST /api/projects/:id/reviews', 'Request AI review'],
        ['GET /api/projects/:id/reviews/:rId', 'Get review with findings'],
        ['PUT /api/projects/:id/reviews/:rId/findings/:fId', 'Accept/reject finding'],
        ['GET /api/projects/:id/versions', 'List architecture versions'],
        ['GET /api/projects/:id/versions/compare', 'Compare two versions'],
        ['POST /api/projects/:id/versions/:vId/rollback', 'Rollback to version'],
        ['POST /api/projects/:id/deployments', 'Initiate deployment'],
        ['POST /api/projects/:id/deployments/:dId/approve', 'Approve deployment'],
        ['POST /api/projects/:id/deployments/:dId/rollback', 'Rollback deployment'],
        ['POST /api/projects/:id/conversations', 'Start chat conversation'],
        ['POST /api/projects/:id/conversations/:cId/messages', 'Send chat message'],
        ['GET /api/projects/:id/conversations/:cId/messages', 'Get chat history'],
        ['GET /api/projects/:id/diagrams/:type', 'Get diagram'],
        ['GET /api/projects/:id/workflow-status', 'Get generation progress'],
    ]
    add_table(doc, ['Endpoint', 'Purpose'], api_data)

    heading2(doc, 'Backend Folder Structure')
    add_folder_structure(doc, [
        "backend/src/",
        "├── controllers/",
        "│   ├── auth.controller.ts        # Login, register, token refresh",
        "│   ├── project.controller.ts      # Project CRUD operations",
        "│   ├── document.controller.ts     # File upload & processing",
        "│   ├── architecture.controller.ts # Architecture generation & retrieval",
        "│   ├── terraform.controller.ts    # Terraform code & validation",
        "│   ├── review.controller.ts       # AI review & findings",
        "│   ├── version.controller.ts      # Version history & comparison",
        "│   ├── deployment.controller.ts   # Deploy, approve, rollback",
        "│   ├── chat.controller.ts         # Conversations & messages",
        "│   ├── diagram.controller.ts      # Diagram generation & export",
        "│   └── cloud.controller.ts        # Cloud service catalog",
        "├── services/",
        "│   ├── auth.service.ts            # JWT, password hashing, sessions",
        "│   ├── project.service.ts         # Project business logic",
        "│   ├── document.service.ts        # File parsing, text extraction",
        "│   ├── architecture.service.ts    # Architecture storage & manipulation",
        "│   ├── terraform.service.ts       # Terraform code management",
        "│   ├── review.service.ts          # Review triggers & findings tracking",
        "│   ├── version.service.ts         # Snapshots, diffs, rollback",
        "│   ├── deployment.service.ts      # Terraform execution, status",
        "│   ├── chat.service.ts            # Message routing to agents",
        "│   ├── notification.service.ts    # WebSocket & email notifications",
        "│   └── audit.service.ts           # Immutable audit log recording",
        "├── models/",
        "│   ├── user.model.ts              # Users table",
        "│   ├── organization.model.ts      # Organizations table",
        "│   ├── project.model.ts           # Projects table",
        "│   ├── document.model.ts          # Uploaded documents table",
        "│   ├── requirement.model.ts       # Extracted requirements table",
        "│   ├── architecture-version.model.ts  # Version snapshots table",
        "│   ├── review.model.ts            # Review results table",
        "│   ├── review-finding.model.ts    # Individual findings table",
        "│   ├── deployment.model.ts        # Deployment records table",
        "│   ├── conversation.model.ts      # Chat conversations table",
        "│   ├── message.model.ts           # Chat messages table",
        "│   ├── knowledge.model.ts         # RAG knowledge base table",
        "│   ├── agent-execution.model.ts   # Agent run logs table",
        "│   └── audit-log.model.ts         # Audit trail table",
        "├── middleware/",
        "│   ├── auth.middleware.ts          # JWT verification",
        "│   ├── rbac.middleware.ts          # Role-based access control",
        "│   ├── validation.middleware.ts    # Request body validation",
        "│   ├── rate-limit.middleware.ts    # API rate limiting",
        "│   └── error-handler.middleware.ts # Global error handler",
        "├── routes/                         # Route files mapping URLs to controllers",
        "├── migrations/                     # All database migration files",
        "├── websocket/                      # WebSocket event handlers",
        "└── utils/                          # JWT helpers, logger, S3 client",
    ])

    add_hr(doc)

    # ===== PHASE 3 =====
    heading1(doc, 'Phase 3 — AI Agents')
    body(doc,
        'Build 16 specialized AI agents that collaborate through an Orchestrator. Each agent handles one '
        'domain — document processing, requirements, knowledge retrieval, architecture design, database, '
        'APIs, cloud mapping, Terraform, security, cost, review, validation, documentation, versioning, '
        'and deployment. Agents share context through a common workspace and communicate via the '
        'Orchestrator which manages execution order, dependencies, and error recovery.'
    )

    heading2(doc, 'What to do')
    bullet(doc, 'Build a base agent class with common LLM calling logic, prompt management, and error handling.')
    bullet(doc, 'Build the Orchestrator Agent that manages the full workflow pipeline and handles retries/failures.')
    bullet(doc, 'Implement a shared workspace (context object) that accumulates all agent outputs.')
    bullet(doc, 'Build each of the 16 agents listed below with their specific prompts and processing logic.')
    bullet(doc, 'Set up vector database (Pinecone/Weaviate) for RAG knowledge storage.')
    bullet(doc, 'Write prompt templates for each agent stored in a prompts/ directory.')
    bullet(doc, 'Build agent execution logging for monitoring and debugging.')

    heading2(doc, 'Agents')
    agents_data = [
        ['Orchestrator', 'Manages workflow order, invokes agents, handles errors and retries.'],
        ['Document Processing', 'Parses uploaded docs, extracts structured text, identifies sections/tables.'],
        ['Requirement', 'Extracts and categorizes requirements (functional, non-functional, constraints).'],
        ['RAG', 'Retrieves relevant enterprise knowledge from vector DB using similarity search.'],
        ['Architecture', 'Generates high-level system design — components, relationships, patterns.'],
        ['Database', 'Designs schemas, tables, relationships, indexes, selects DB technology.'],
        ['API', 'Designs REST endpoints, request/response schemas, produces OpenAPI spec.'],
        ['Cloud Mapping', 'Maps components to cloud services (AWS/Azure/GCP), estimates costs.'],
        ['Terraform', 'Generates production-ready Terraform modules from cloud mappings.'],
        ['Security', 'Validates architecture against security best practices and compliance rules.'],
        ['Cost Optimization', 'Estimates cloud costs, identifies savings, suggests right-sizing.'],
        ['Architecture Review', 'Scores architecture on completeness, consistency, security, scalability.'],
        ['Validation', 'Checks Terraform syntax, schema validity, cross-reference consistency.'],
        ['Documentation', 'Generates architecture docs, API docs, deployment guides, ADRs.'],
        ['Versioning', 'Creates immutable snapshots, generates change diffs between versions.'],
        ['Deployment', 'Executes Terraform init/plan/apply, monitors deployment, handles rollback.'],
    ]
    add_table(doc, ['Agent', 'Responsibility'], agents_data)

    heading2(doc, 'Agent Execution Flow')
    flow_steps = [
        "User uploads document",
        "Document Processing Agent → structures the content",
        "Requirement Agent → extracts categorized requirements",
        "RAG Agent → retrieves enterprise knowledge",
        "Architecture Agent → generates system design",
        "Database Agent → creates schema  |  API Agent → designs endpoints  (parallel)",
        "Cloud Mapping Agent → selects cloud services, estimates costs",
        "Terraform Agent → generates infrastructure code",
        "Security Agent → validates security  |  Validation Agent → checks correctness  (parallel)",
        "Cost Optimization Agent → analyzes pricing",
        "Architecture Review Agent → scores and produces findings",
        "Documentation Agent → generates docs",
        "Versioning Agent → creates version snapshot",
        "User receives complete architecture",
    ]
    for i, step in enumerate(flow_steps):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.3)
        run_num = p.add_run(f"{i+1}.  ")
        run_num.bold = True
        run_num.font.size = Pt(10)
        run_num.font.name = 'Arial'
        run_text = p.add_run(step)
        run_text.font.size = Pt(10)
        run_text.font.name = 'Arial'
        if i < len(flow_steps) - 1:
            arrow = doc.add_paragraph()
            arrow.paragraph_format.space_before = Pt(0)
            arrow.paragraph_format.space_after = Pt(0)
            arrow.paragraph_format.left_indent = Inches(0.6)
            run_a = arrow.add_run('↓')
            run_a.font.size = Pt(9)
            run_a.font.color.rgb = RGBColor(136, 136, 136)

    heading2(doc, 'Agents Folder Structure')
    add_folder_structure(doc, [
        "agents/",
        "├── src/",
        "│   ├── common/",
        "│   │   ├── base.agent.ts          # Base class — LLM calls, retries, logging",
        "│   │   ├── agent.context.ts        # Shared workspace between agents",
        "│   │   ├── llm.client.ts           # LLM API client (OpenAI/Groq/Anthropic)",
        "│   │   └── prompt.manager.ts       # Loads and fills prompt templates",
        "│   ├── orchestrator/",
        "│   │   ├── orchestrator.agent.ts    # Workflow manager, agent sequencing",
        "│   │   ├── workflow.manager.ts      # Execution order, dependency graph",
        "│   │   └── agent.registry.ts        # Registers all available agents",
        "│   ├── document-processing/",
        "│   │   └── document.agent.ts        # Text extraction, structure parsing",
        "│   ├── requirement/",
        "│   │   └── requirement.agent.ts     # Requirement identification & categorization",
        "│   ├── rag/",
        "│   │   ├── rag.agent.ts             # Query formulation, knowledge retrieval",
        "│   │   ├── embeddings.service.ts    # Chunk text and generate embeddings",
        "│   │   └── vector.store.ts          # Pinecone/Weaviate client",
        "│   ├── architecture/",
        "│   │   └── architecture.agent.ts    # System design, component selection",
        "│   ├── database/",
        "│   │   └── database.agent.ts        # Schema design, tech selection",
        "│   ├── api/",
        "│   │   └── api.agent.ts             # Endpoint design, OpenAPI generation",
        "│   ├── cloud-mapping/",
        "│   │   ├── cloud-mapping.agent.ts   # Service selection, cost estimation",
        "│   │   └── providers/",
        "│   │       ├── aws.catalog.ts       # AWS service catalog & pricing",
        "│   │       ├── azure.catalog.ts     # Azure service catalog & pricing",
        "│   │       └── gcp.catalog.ts       # GCP service catalog & pricing",
        "│   ├── terraform/",
        "│   │   ├── terraform.agent.ts       # Terraform code generation",
        "│   │   └── generators/",
        "│   │       ├── aws/                 # AWS-specific Terraform modules",
        "│   │       ├── azure/               # Azure-specific Terraform modules",
        "│   │       └── gcp/                 # GCP-specific Terraform modules",
        "│   ├── security/",
        "│   │   ├── security.agent.ts        # Security validation & compliance",
        "│   │   └── policies/                # Compliance rules (SOC2, HIPAA, GDPR)",
        "│   ├── cost-optimization/",
        "│   │   └── cost.agent.ts            # Cost analysis & optimization",
        "│   ├── review/",
        "│   │   └── review.agent.ts          # Quality scoring & recommendations",
        "│   ├── validation/",
        "│   │   └── validation.agent.ts      # Syntax & consistency checks",
        "│   ├── documentation/",
        "│   │   ├── documentation.agent.ts   # Doc generation",
        "│   │   └── templates/               # Documentation templates",
        "│   ├── versioning/",
        "│   │   ├── versioning.agent.ts      # Snapshot creation",
        "│   │   └── diff.engine.ts           # Version comparison engine",
        "│   ├── deployment/",
        "│   │   └── deployment.agent.ts      # Terraform execution & monitoring",
        "│   └── types/                        # Agent type definitions",
        "├── prompts/                          # All prompt templates per agent",
        "├── tests/                            # Agent unit tests",
        "└── package.json",
    ])

    add_hr(doc)

    # End of document
    p = doc.add_paragraph()
    run = p.add_run('End of Document')
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = 'Arial'

    doc.save(OUTPUT)
    print(f"DOCX generated: {OUTPUT}")


if __name__ == "__main__":
    build()
